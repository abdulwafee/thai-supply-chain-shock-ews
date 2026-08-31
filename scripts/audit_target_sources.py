"""Task A1 audit: verify real, inspected data coverage for the four candidate
Industry Stress Index target components (docs/target_definition.md §1).

This is a **quality-gated revision**. The first version of this script made
two claims that did not hold up to scrutiny and are fixed here:

1. It reported "12/12 industries covered" without distinguishing full from
   partial coverage (IND-03 and IND-12 are each missing one constituent TSIC
   division) — see `IndustryCoverage` below, which now reports all three
   counts (any / full / partial) separately, never collapsed into one number.
2. It stored per-division missingness in a plain `dict`, which would have
   silently overwritten an earlier row's data if a TSIC division label ever
   appeared twice in a workbook. No duplicate exists in the two files
   inspected (verified explicitly below, not assumed), but the script now
   raises `DuplicateDivisionError` if one ever does, rather than silently
   keeping only the last-seen row.

This script is reproducible evidence, not a pipeline component — it does not
belong to src/thai_supply_chain_ews/ because it performs one-off source
verification (Phase 1 Group A), not recurring ingestion logic (that is
data/ingest.py, still a stub, and stays a stub after this script — see the
module-level restriction against turning this into the production pipeline).

Verifying `source access` is not the same claim as `target-component
approval` — this script computes evidence for the latter but does not itself
grant it. See docs/target_definition.md §1.4 and configs/targets.yaml's
`source_verified` / `target_approved` fields for where that decision is
actually recorded.

No absolute paths: everything is resolved from this script's own location
(project root = two directories up).

Usage:
    python scripts/audit_target_sources.py
"""

from __future__ import annotations

import hashlib
import json
import sys
from dataclasses import dataclass, field
from pathlib import Path

try:
    import openpyxl
except ImportError:  # pragma: no cover
    print("openpyxl is required: pip install openpyxl", file=sys.stderr)
    raise

try:
    import yaml
except ImportError:  # pragma: no cover
    print("pyyaml is required: pip install pyyaml", file=sys.stderr)
    raise

PROJECT_ROOT = Path(__file__).resolve().parents[1]
RAW_DIR = PROJECT_ROOT / "data" / "raw"
TARGETS_CONFIG_PATH = PROJECT_ROOT / "configs" / "targets.yaml"

# --- shared parsing primitives now live in the installable package ----------
# Task B1 moved the low-level OIE parsing logic into
# `thai_supply_chain_ews.data.oie` so there is ONE production implementation.
# This script imports it rather than maintaining a second, divergent copy.
# The dependency runs one way only: the package never imports this script.
sys.path.insert(0, str(PROJECT_ROOT / "src"))
from thai_supply_chain_ews.data import oie as _oie  # noqa: E402

DuplicateDivisionError = _oie.DuplicateDivisionError
TSIC_ROW_PATTERN = _oie.TSIC_ROW_PATTERN
THAI_MONTH_TO_NUMBER = _oie.THAI_MONTH_TO_NUMBER
THAI_MONTH_ABBREVIATIONS = _oie.THAI_MONTH_ABBREVIATIONS
BUDDHIST_TO_GREGORIAN_YEAR_OFFSET = _oie.BUDDHIST_TO_GREGORIAN_YEAR_OFFSET
MonthlyDateAudit = _oie.MonthlyDateAudit
find_tsic_label_column = _oie.find_tsic_label_column
find_first_month_column = _oie.find_first_month_column
parse_monthly_dates = _oie.parse_monthly_dates


class EligibilityConstantMismatchError(ValueError):
    """This script's hard-coded eligibility constants no longer match
    configs/targets.yaml's eligibility_criteria block. Point 9 of the
    quality-gate review: don't let the two silently drift apart.
    """


# --- eligibility criteria, cross-checked against configs/targets.yaml at
# runtime (see check_eligibility_constants_match_config) rather than
# imported directly — this script must run standalone, but "standalone"
# must not mean "silently divergent" ------------------------------------
MIN_TIME_COVERAGE_MONTHS = 60
MIN_INDUSTRY_COVERAGE = 11  # of 12, "any coverage" — the ORIGINAL, coarser
# eligibility bar. Passing this is necessary
# but not sufficient for target-component
# approval — see FULL vs PARTIAL below.
MAX_MISSINGNESS_PCT = 10.0
PREFERRED_HISTORY_MONTHS = 120  # not an eligibility gate, a target-construction
# preference — see docs/task_a1_report.md §7

# --- FROZEN DRAFT TAXONOMY — do not "fix" to match production ---------------
# This is the DRAFT mapping used by Task A1/A2. It is deliberately NOT the
# production taxonomy. Production reads configs/industry_mapping.yaml via
# `thai_supply_chain_ews.data.taxonomy` (see docs/architecture/decision_log.md
# AD-R12).
#
# It is frozen here because the preserved bridge-validation evidence
# (oie_overlap_bridge_v1 / v2) was computed against these definitions, and
# those artifacts are pinned by SHA-256. Changing this constant would silently
# invalidate that evidence rather than superseding it honestly.
#
# Known differences from the production taxonomy, all corrected in B1:
#   * TSIC 12 is absent here entirely — a real gap; production assigns it to IND-01.
#   * IND-03 includes TSIC 18 and IND-12 includes TSIC 33; neither is reported
#     by the current OIE files, so production excludes both.
DRAFT_INDUSTRY_TSIC_DIVISIONS_A1A2_FROZEN = {
    "IND-01": [10, 11],
    "IND-02": [13, 14, 15],
    "IND-03": [16, 17, 18],
    "IND-04": [19],
    "IND-05": [20, 21],
    "IND-06": [22],
    "IND-07": [23],
    "IND-08": [24, 25],
    "IND-09": [26, 27],
    "IND-10": [28],
    "IND-11": [29, 30],
    "IND-12": [31, 32, 33],
}

# Backwards-compatible alias. scripts/validate_oie_overlap.py (v1) and
# validate_oie_overlap_v2.py read this name and are checksum-pinned, so it
# must keep resolving to the frozen draft above.
INDUSTRY_TSIC_DIVISIONS = DRAFT_INDUSTRY_TSIC_DIVISIONS_A1A2_FROZEN


def check_eligibility_constants_match_config() -> None:
    """Point 9: assert this script's constants match configs/targets.yaml
    rather than silently duplicating them with no cross-check.
    """
    if not TARGETS_CONFIG_PATH.is_file():
        raise FileNotFoundError(f"{TARGETS_CONFIG_PATH} not found — cannot cross-check constants")
    config = yaml.safe_load(TARGETS_CONFIG_PATH.read_text(encoding="utf-8"))
    criteria = config.get("eligibility_criteria", {})
    mismatches = []
    checks = [
        ("minimum_time_coverage_months", MIN_TIME_COVERAGE_MONTHS),
        ("minimum_industry_coverage", MIN_INDUSTRY_COVERAGE),
        ("maximum_missingness_pct", MAX_MISSINGNESS_PCT),
        ("preferred_history_months", PREFERRED_HISTORY_MONTHS),
    ]
    for key, script_value in checks:
        config_value = criteria.get(key)
        if config_value != script_value:
            mismatches.append(f"{key}: script has {script_value!r}, config has {config_value!r}")
    if mismatches:
        raise EligibilityConstantMismatchError(
            "This script's eligibility constants no longer match configs/targets.yaml:\n"
            + "\n".join(mismatches)
        )


@dataclass
class IndustryCoverage:
    """Full vs. partial vs. none — never collapsed into one "N/12" figure."""

    fully_covered: list[str] = field(default_factory=list)
    partially_covered: dict[str, list[int]] = field(default_factory=dict)
    not_covered: list[str] = field(default_factory=list)

    @property
    def any_coverage_count(self) -> int:
        return len(self.fully_covered) + len(self.partially_covered)

    @property
    def full_coverage_count(self) -> int:
        return len(self.fully_covered)

    @property
    def partial_coverage_count(self) -> int:
        return len(self.partially_covered)


@dataclass
class OieAuditResult:
    file_path: Path
    total_row_count: int
    dates: MonthlyDateAudit
    tsic_divisions_found: list[int] = field(default_factory=list)
    tsic_divisions_missing: list[int] = field(default_factory=list)
    per_division_missingness_pct: dict[int, float] = field(default_factory=dict)
    per_division_weight: dict[int, float] = field(default_factory=dict)
    weight_sum: float = 0.0
    industry_coverage: IndustryCoverage = field(default_factory=IndustryCoverage)


# find_tsic_label_column / find_first_month_column / parse_monthly_dates and
# MonthlyDateAudit are imported from thai_supply_chain_ews.data.oie at the top of
# this module (Task B1 de-duplication). They are re-exported under their original
# names so this script's public API is unchanged.


def audit_oie_file(path: Path) -> OieAuditResult:
    if not path.is_file():
        raise FileNotFoundError(
            f"{path} not found. This audit does not fabricate results for a "
            "file it cannot read — re-run the retrieval step first."
        )
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]  # "รายเดือน" (monthly) is always the first sheet
    rows = list(ws.iter_rows(min_row=1, max_row=ws.max_row, values_only=True))

    label_col = find_tsic_label_column(rows)
    month_col_start = find_first_month_column(rows, label_col)
    date_audit = parse_monthly_dates(rows, month_col_start)
    month_count = len(date_audit.dates)

    # weight column is immediately to the right of the label column in both
    # files inspected (verified: MPI label_col=2 -> weight_col=3;
    # CapU label_col=0 -> weight_col=1)
    weight_col = label_col + 1

    found_divisions: list[int] = []
    seen_divisions: set[int] = set()
    missingness: dict[int, float] = {}
    weights: dict[int, float] = {}
    for row_idx, row in enumerate(rows, start=1):
        label = row[label_col] if label_col < len(row) else None
        if not (isinstance(label, str) and TSIC_ROW_PATTERN.match(label)):
            continue
        division = int(TSIC_ROW_PATTERN.match(label).group(1))

        if division in seen_divisions:
            raise DuplicateDivisionError(
                f"TSIC division {division} appears more than once in {path.name} "
                f"(second occurrence at row {row_idx}) — refusing to silently "
                "overwrite the first row's data. Resolve which row is the "
                "official division-total before re-running this audit."
            )
        seen_divisions.add(division)
        found_divisions.append(division)

        values = row[month_col_start : month_col_start + month_count]
        n_missing = sum(1 for v in values if v is None or v == "")
        missingness[division] = 100.0 * n_missing / month_count if month_count else 100.0

        weight_val = row[weight_col] if weight_col < len(row) else None
        try:
            weights[division] = float(weight_val)
        except (TypeError, ValueError):
            weights[division] = float("nan")

    all_expected = set(range(10, 34))
    missing_divisions = sorted(all_expected - set(found_divisions))

    coverage = IndustryCoverage()
    for industry_id, divisions in INDUSTRY_TSIC_DIVISIONS.items():
        present = [d for d in divisions if d in found_divisions]
        if len(present) == len(divisions):
            coverage.fully_covered.append(industry_id)
        elif present:
            coverage.partially_covered[industry_id] = sorted(set(divisions) - set(present))
        else:
            coverage.not_covered.append(industry_id)

    return OieAuditResult(
        file_path=path,
        total_row_count=len(rows),
        dates=date_audit,
        tsic_divisions_found=sorted(found_divisions),
        tsic_divisions_missing=missing_divisions,
        per_division_missingness_pct=missingness,
        per_division_weight=weights,
        weight_sum=sum(v for v in weights.values() if v == v),  # NaN-safe sum
        industry_coverage=coverage,
    )


def sha256_of(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def resolve_retrieval_path(source_id: str) -> Path:
    """Resolve the file to audit from the retrieval manifest, not a
    hard-coded date directory (point 9 of the quality-gate review).

    Falls back to the newest dated subdirectory under data/raw/<source_id>/
    if no manifest entry is found, but a manifest entry is preferred since
    it is the actual record of what was verified and when.
    """
    manifest_path = RAW_DIR / "_manifests" / f"{source_id}_manifest.jsonl"
    if manifest_path.is_file():
        lines = manifest_path.read_text(encoding="utf-8").strip().splitlines()
        if lines:
            last_entry = json.loads(lines[-1])
            return PROJECT_ROOT / last_entry["file_path"]

    source_dir = RAW_DIR / source_id
    if not source_dir.is_dir():
        raise FileNotFoundError(f"No manifest and no directory for {source_id} under {RAW_DIR}")
    dated_subdirs = sorted(p for p in source_dir.iterdir() if p.is_dir())
    if not dated_subdirs:
        raise FileNotFoundError(f"No dated retrieval subdirectory found under {source_dir}")
    newest = dated_subdirs[-1]
    files = list(newest.glob("*.xlsx"))
    if not files:
        raise FileNotFoundError(f"No .xlsx file found under {newest}")
    return files[0]


def print_report(name: str, result: OieAuditResult) -> None:
    print(f"\n=== {name} ===")
    print(f"file: {result.file_path.relative_to(PROJECT_ROOT)}")
    print(f"sha256: {sha256_of(result.file_path)}")
    print(f"total rows in sheet: {result.total_row_count}")

    d = result.dates
    print(
        f"parsed calendar months: {len(d.dates)}  "
        f"(eligibility minimum: {MIN_TIME_COVERAGE_MONTHS}, preferred: {PREFERRED_HISTORY_MONTHS})"
    )
    print(
        f"  earliest: {d.earliest.isoformat()}   latest: {d.latest.isoformat()}"
        f"{' (preliminary)' if d.preliminary_flags[-1] else ''}"
    )
    print(
        f"  chronological: {d.is_chronological}   duplicate months: {d.has_duplicates}   "
        f"gap months: {d.gap_months or 'none'}"
    )
    date_integrity_ok = (
        len(d.dates) >= MIN_TIME_COVERAGE_MONTHS
        and d.is_chronological
        and not d.has_duplicates
        and not d.gap_months
    )
    print(f"  -> {'PASS' if date_integrity_ok else 'FAIL'} on date-coverage integrity")

    print(
        f"TSIC divisions found ({len(result.tsic_divisions_found)}): {result.tsic_divisions_found}"
    )
    print(f"TSIC divisions missing: {result.tsic_divisions_missing}")
    print(
        f"sum of division-level weights: {result.weight_sum:.4f}  "
        "(0-value or NaN entries excluded; ~100 confirms these divisions are the "
        "survey's full weighted universe, not a subset with unaccounted-for weight)"
    )

    cov = result.industry_coverage
    print(
        f"industries — any coverage: {cov.any_coverage_count}/12 "
        f"(eligibility minimum: {MIN_INDUSTRY_COVERAGE})  "
        f"| FULL coverage: {cov.full_coverage_count}/12  "
        f"| PARTIAL coverage: {cov.partial_coverage_count}/12"
    )
    print(f"  -> fully covered: {sorted(cov.fully_covered)}")
    print(f"  -> partially covered (missing division listed): {cov.partially_covered}")
    print(f"  -> not covered at all: {cov.not_covered or 'none'}")
    print(
        f"  -> {'PASS' if cov.any_coverage_count >= MIN_INDUSTRY_COVERAGE else 'FAIL'} "
        "on the ORIGINAL any-coverage eligibility bar (this is NOT the same claim as "
        "'ready for final 12-industry aggregation' — see docs/task_a1_report.md §4)"
    )

    worst_missingness = max(result.per_division_missingness_pct.values(), default=100.0)
    print(
        f"worst per-division missingness: {worst_missingness:.2f}%  "
        f"(eligibility maximum: {MAX_MISSINGNESS_PCT}%)"
    )
    print(
        f"  -> {'PASS' if worst_missingness <= MAX_MISSINGNESS_PCT else 'FAIL'} "
        "on maximum missingness"
    )


def main() -> None:
    check_eligibility_constants_match_config()

    mpi_path = resolve_retrieval_path("OIE_MPI")
    capu_path = resolve_retrieval_path("OIE_CAPU")

    results: dict[str, OieAuditResult] = {}
    for name, path in [("OIE_MPI", mpi_path), ("OIE_CAPU", capu_path)]:
        try:
            result = audit_oie_file(path)
            results[name] = result
            print_report(name, result)
        except (FileNotFoundError, DuplicateDivisionError) as exc:
            print(f"\n=== {name} ===\nBLOCKED: {exc}")

    output = {}
    for name, r in results.items():
        d = r.dates
        output[name] = {
            "file": str(r.file_path.relative_to(PROJECT_ROOT)),
            "sha256": sha256_of(r.file_path),
            "dates": {
                "parsed_months": [dt.isoformat() for dt in d.dates],
                "preliminary_flags": d.preliminary_flags,
                "earliest": d.earliest.isoformat(),
                "latest": d.latest.isoformat(),
                "is_chronological": d.is_chronological,
                "has_duplicate_months": d.has_duplicates,
                "gap_months": d.gap_months,
                "count": len(d.dates),
            },
            "tsic_divisions_found": r.tsic_divisions_found,
            "tsic_divisions_missing": r.tsic_divisions_missing,
            "per_division_missingness_pct": r.per_division_missingness_pct,
            "per_division_weight": r.per_division_weight,
            "weight_sum": r.weight_sum,
            "industry_coverage": {
                "fully_covered": sorted(r.industry_coverage.fully_covered),
                "partially_covered": r.industry_coverage.partially_covered,
                "not_covered": r.industry_coverage.not_covered,
                "any_coverage_count": r.industry_coverage.any_coverage_count,
                "full_coverage_count": r.industry_coverage.full_coverage_count,
                "partial_coverage_count": r.industry_coverage.partial_coverage_count,
            },
        }
    report_path = PROJECT_ROOT / "docs" / "task_a1_audit_output.json"
    report_path.write_text(json.dumps(output, indent=2, ensure_ascii=False), encoding="utf-8")
    print(f"\nMachine-readable output written to {report_path.relative_to(PROJECT_ROOT)}")
    print(
        "\nMOC_EXPORT and NSO_PPI are NOT audited by this script: both require "
        "interactive query-parameter discovery that this script does not guess. "
        "See docs/task_a1_report.md for what was verified manually.\n"
        "\nNOTE: this script computes SOURCE-ACCESS evidence only. It does not "
        "grant TARGET-COMPONENT approval — see configs/targets.yaml's "
        "source_verified vs. target_approved fields and docs/target_definition.md §1.4."
    )


if __name__ == "__main__":
    main()
