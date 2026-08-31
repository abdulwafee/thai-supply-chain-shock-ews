"""Low-level OIE workbook parsing — the single production implementation.

This module owns the parsing primitives that were originally written inside
`scripts/audit_target_sources.py` during Phase 1 Group A. They now live here,
in the installable package, and the audit scripts import them from here rather
than keeping a second, divergent copy. The dependency direction is one-way:
production code never imports the audit scripts.

What the OIE workbooks actually look like (verified against both real
current-edition files, not assumed):

* A hierarchical sheet: Total -> Division ("TSIC : NN <thai name>") -> Group
  ("TSIC : NNNN") -> Class -> Product. Only the two-digit division rows are
  aggregation inputs.
* The TSIC label column differs between files (index 2 in the MPI workbook,
  index 0 in the Capacity Utilization workbook), so it is detected, not fixed.
* The column immediately after the label is the official value-added WEIGHT,
  not a month. Confusing the two was a real bug once; `find_first_month_column`
  therefore reads the header text, never a data row's numeric run.
* Header dates are Buddhist-Era years printed once per 12-column block, above
  a row of Thai month abbreviations. A trailing "*" on a month abbreviation
  marks a preliminary observation.
"""

from __future__ import annotations

import hashlib
import re
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path

try:
    import openpyxl
except ImportError as exc:  # pragma: no cover
    raise ImportError("openpyxl is required to parse OIE workbooks") from exc

__all__ = [
    "BUDDHIST_TO_GREGORIAN_YEAR_OFFSET",
    "THAI_MONTH_ABBREVIATIONS",
    "THAI_MONTH_TO_NUMBER",
    "TSIC_ROW_PATTERN",
    "DivisionRow",
    "DuplicateDivisionError",
    "MonthlyDateAudit",
    "OieParseError",
    "ParsedWorkbook",
    "buddhist_to_gregorian_year",
    "find_first_month_column",
    "find_tsic_label_column",
    "parse_monthly_dates",
    "parse_workbook",
    "read_workbook_rows",
    "sha256_of_file",
]

TSIC_ROW_PATTERN = re.compile(r"^TSIC\s*:\s*(\d\d)\s")

THAI_MONTH_TO_NUMBER = {
    "ม.ค.": 1, "ก.พ.": 2, "มี.ค.": 3, "เม.ย.": 4,
    "พ.ค.": 5, "มิ.ย.": 6, "ก.ค.": 7, "ส.ค.": 8,
    "ก.ย.": 9, "ต.ค.": 10, "พ.ย.": 11, "ธ.ค.": 12,
}
THAI_MONTH_ABBREVIATIONS = frozenset(THAI_MONTH_TO_NUMBER)
BUDDHIST_TO_GREGORIAN_YEAR_OFFSET = 543

PRELIMINARY_MARKER = "*"


class OieParseError(ValueError):
    """The workbook did not match the verified OIE layout."""


class DuplicateDivisionError(OieParseError):
    """A TSIC division label appeared more than once in one workbook.

    Refuses to silently keep only one of the rows — a dict-based store would
    have overwritten the earlier one without any signal.
    """


def buddhist_to_gregorian_year(be_year: int) -> int:
    """2564 -> 2021. Rejects a year that is obviously already Gregorian, since
    silently subtracting 543 from a Gregorian year would shift the whole series
    by five centuries.
    """
    if be_year < 2400:
        raise OieParseError(
            f"{be_year} does not look like a Buddhist-Era year (expected >= 2400); "
            "refusing to convert."
        )
    return be_year - BUDDHIST_TO_GREGORIAN_YEAR_OFFSET


def sha256_of_file(path: Path) -> str:
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def read_workbook_rows(path: Path) -> list[tuple]:
    """Read the first worksheet as a list of value-only row tuples."""
    path = Path(path)
    if not path.is_file():
        raise FileNotFoundError(f"{path} not found")
    workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
    try:
        sheet = workbook[workbook.sheetnames[0]]
        return list(sheet.iter_rows(min_row=1, max_row=sheet.max_row, values_only=True))
    finally:
        workbook.close()


def find_tsic_label_column(rows: list[tuple]) -> int:
    """Detect which column carries the "TSIC : NN" labels."""
    for row in rows[:50]:
        for col_idx, value in enumerate(row[:6]):
            if isinstance(value, str) and TSIC_ROW_PATTERN.match(value):
                return col_idx
    raise OieParseError("Could not locate a TSIC-labeled row in the first 50 rows")


def find_first_month_column(rows: list[tuple], label_col: int) -> int:
    """Locate the first monthly data column from the HEADER text.

    Deliberately not inferred from a data row's first numeric cell: the weight
    column sits immediately after the label and is also numeric, so that
    approach silently returns the weight column.
    """
    del label_col  # kept for call-signature compatibility with the audit scripts
    for row in rows[:10]:
        for col_idx, value in enumerate(row):
            if isinstance(value, str) and value.rstrip(PRELIMINARY_MARKER) in (
                THAI_MONTH_ABBREVIATIONS
            ):
                return col_idx
    raise OieParseError("Could not locate the first monthly data column from the header row")


@dataclass
class MonthlyDateAudit:
    """Explicitly parsed calendar coverage — never just a column count."""

    dates: list[date]
    preliminary_flags: list[bool]
    earliest: date
    latest: date
    is_chronological: bool
    has_duplicates: bool
    gap_months: list[str]


def _next_month(d: date) -> date:
    return date(d.year + (1 if d.month == 12 else 0), 1 if d.month == 12 else d.month + 1, 1)


def parse_monthly_dates(rows: list[tuple], month_col_start: int) -> MonthlyDateAudit:
    """Parse the Buddhist-year row + Thai-month row into Gregorian dates.

    The year is printed once per 12-column block, so the most recent non-null
    year cell carries forward across that block. A trailing "*" marks the
    observation preliminary and is recorded, not stripped away and forgotten.
    """
    month_row_idx = None
    for i, row in enumerate(rows[:10]):
        candidates = [
            str(v).rstrip(PRELIMINARY_MARKER)
            for v in row[month_col_start : month_col_start + 5]
            if isinstance(v, str)
        ]
        if any(c in THAI_MONTH_ABBREVIATIONS for c in candidates):
            month_row_idx = i
            break
    if month_row_idx is None:
        raise OieParseError("Could not locate the month-abbreviation header row")
    if month_row_idx == 0:
        raise OieParseError("Month header row has no year row above it")

    year_row = rows[month_row_idx - 1]
    month_row = rows[month_row_idx]

    dates: list[date] = []
    preliminary_flags: list[bool] = []
    current_be_year: int | None = None
    for col_idx in range(month_col_start, len(month_row)):
        year_cell = year_row[col_idx] if col_idx < len(year_row) else None
        if isinstance(year_cell, (int, float)) and not isinstance(year_cell, bool):
            current_be_year = int(year_cell)

        month_cell = month_row[col_idx] if col_idx < len(month_row) else None
        if not isinstance(month_cell, str):
            break  # trailing "% change" columns or end of row
        stripped = month_cell.rstrip(PRELIMINARY_MARKER)
        if stripped not in THAI_MONTH_ABBREVIATIONS:
            break
        if current_be_year is None:
            raise OieParseError(f"Month column {col_idx} has no preceding Buddhist year")
        dates.append(
            date(buddhist_to_gregorian_year(current_be_year), THAI_MONTH_TO_NUMBER[stripped], 1)
        )
        preliminary_flags.append(month_cell.endswith(PRELIMINARY_MARKER))

    if not dates:
        raise OieParseError("Parsed zero calendar dates — header layout assumption failed")

    is_chronological = dates == sorted(dates)
    has_duplicates = len(dates) != len(set(dates))

    gap_months: list[str] = []
    if is_chronological and not has_duplicates:
        expected = dates[0]
        for d in dates:
            while expected < d:
                gap_months.append(f"{expected.year:04d}-{expected.month:02d}")
                expected = _next_month(expected)
            expected = _next_month(d)

    return MonthlyDateAudit(
        dates=dates,
        preliminary_flags=preliminary_flags,
        earliest=dates[0],
        latest=dates[-1],
        is_chronological=is_chronological,
        has_duplicates=has_duplicates,
        gap_months=gap_months,
    )


@dataclass
class DivisionRow:
    """One two-digit TSIC division-total row."""

    tsic_division: int
    tsic_label: str
    source_weight: float
    values_by_month: dict[date, float]
    non_numeric_months: list[date] = field(default_factory=list)


@dataclass
class ParsedWorkbook:
    file_path: Path
    sha256: str
    dates: MonthlyDateAudit
    divisions: dict[int, DivisionRow]

    @property
    def division_numbers(self) -> list[int]:
        return sorted(self.divisions)

    @property
    def weight_sum(self) -> float:
        return sum(d.source_weight for d in self.divisions.values())


def parse_workbook(path: Path) -> ParsedWorkbook:
    """Parse one OIE workbook into division rows with explicit monthly dates.

    A non-numeric cell is recorded in `non_numeric_months` and simply absent
    from `values_by_month` — never coerced to zero and never imputed. Callers
    decide what a missing value means; this layer does not guess.
    """
    path = Path(path)
    rows = read_workbook_rows(path)
    label_col = find_tsic_label_column(rows)
    month_col_start = find_first_month_column(rows, label_col)
    dates = parse_monthly_dates(rows, month_col_start)
    weight_col = label_col + 1
    month_count = len(dates.dates)

    divisions: dict[int, DivisionRow] = {}
    for row in rows:
        label = row[label_col] if label_col < len(row) else None
        if not isinstance(label, str):
            continue
        match = TSIC_ROW_PATTERN.match(label)
        if not match:
            continue
        division = int(match.group(1))
        if division in divisions:
            raise DuplicateDivisionError(
                f"TSIC division {division} appears more than once in {path.name} — "
                "refusing to silently keep one row."
            )

        raw_weight = row[weight_col] if weight_col < len(row) else None
        if not isinstance(raw_weight, (int, float)) or isinstance(raw_weight, bool):
            raise OieParseError(
                f"TSIC division {division} in {path.name} has a non-numeric source weight "
                f"({raw_weight!r}) — refusing to guess it."
            )
        weight = float(raw_weight)
        if weight <= 0:
            raise OieParseError(
                f"TSIC division {division} in {path.name} has a non-positive source weight "
                f"({weight}) — cannot renormalize."
            )

        cells = row[month_col_start : month_col_start + month_count]
        values: dict[date, float] = {}
        non_numeric: list[date] = []
        for month, cell in zip(dates.dates, cells, strict=False):
            if isinstance(cell, (int, float)) and not isinstance(cell, bool):
                values[month] = float(cell)
            else:
                non_numeric.append(month)

        divisions[division] = DivisionRow(
            tsic_division=division,
            tsic_label=label.strip(),
            source_weight=weight,
            values_by_month=values,
            non_numeric_months=non_numeric,
        )

    if not divisions:
        raise OieParseError(f"No TSIC division-total rows found in {path.name}")

    return ParsedWorkbook(
        file_path=path, sha256=sha256_of_file(path), dates=dates, divisions=divisions
    )
