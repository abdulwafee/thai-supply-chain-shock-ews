"""NESDC Input–Output Table acquisition and parsing — Task C3.

The official Thai I/O table is published by the National Economic and Social
Development Council. This module discovers the download URLs from NESDC's own
pages, retrieves the files, and parses the workbook into the blocks an I/O
analysis needs. Nothing here is hard-coded to an opaque file URL: NESDC migrated
their site during this project's lifetime and the old ``main.php?filename=io_page``
path now serves a WordPress installation, so discovery walks the official
sitemap and download pages instead.

Workbook layout, established by inspection rather than assumed
---------------------------------------------------------------
The published file is NOT a rectangular matrix. It is a SPARSE LONG table with
one row per (ROW, COLUMN) cell and seven columns::

    ROW | COLUMN | PURCHASER | WHOLESALE | RETAIL | TRANSPORT | IMPORT

* ``ROW`` and ``COLUMN`` are three-digit codes.
* Codes ``001``–``180`` are production sectors.
* Row codes ``190`` (total intermediate), ``201``–``204`` (value-added
  components), ``209`` (total value added) and ``210`` (total input) are
  ACCOUNTING TOTALS, not sectors.
* Column codes ``190`` (total intermediate demand), ``301``–``306`` (final
  demand components, ``309`` their subtotal), ``401``–``404`` (imports, entered
  NEGATIVE, ``409`` their subtotal), ``501``–``503`` (stock changes, ``509``
  their subtotal), ``310``/``700`` (total supply) and ``600`` (gross output).

The subtotal codes are the trap: summing ``301``–``309`` double-counts, and
including ``190``/``600``/``700`` inside the intermediate block would fold
totals into the coefficients. This module extracts the intermediate block by
explicit sector-code membership, never by worksheet coordinates.

Valuation
---------
``PURCHASER`` is at PURCHASERS' PRICES and is IMPORT-INCLUSIVE. The margin
columns itemise the trade and transport margins inside that value, and
``IMPORT`` gives the imported content of the same cell — so import-specific
coefficients ARE available from this source and are not inferred.
"""

from __future__ import annotations

import re
import urllib.error
import urllib.request
from dataclasses import dataclass, field
from pathlib import Path

__all__ = [
    "IOSourceError",
    "IOParseError",
    "WafChallengeError",
    "SECTOR_CODES",
    "SECTOR_CODE_COUNT",
    "VALUE_ADDED_ROW_CODES",
    "TOTAL_ROW_CODES",
    "FINAL_DEMAND_COLUMN_CODES",
    "IMPORT_COLUMN_CODES",
    "STOCK_COLUMN_CODES",
    "SUBTOTAL_COLUMN_CODES",
    "GROSS_OUTPUT_COLUMN",
    "TOTAL_INPUT_ROW",
    "TOTAL_VALUE_ADDED_ROW",
    "VALUATION_BASIS",
    "NESDC_SITEMAP_INDEX",
    "IOTable",
    "discover_io_downloads",
    "resolve_download_url",
    "fetch",
    "parse_io_workbook",
    "parse_sector_definitions",
    "NESDC_MEDIA_API",
    "AGGREGATION_LEVELS",
    "WorkbookIdentityError",
    "discover_media_files",
    "validate_workbook_identity",
]

PROJECT_ROOT = Path(__file__).resolve().parents[3]

# --- code universe (validated against the published file) --------------------
SECTOR_CODE_COUNT = 180
SECTOR_CODES: tuple[str, ...] = tuple(f"{i:03d}" for i in range(1, SECTOR_CODE_COUNT + 1))

VALUE_ADDED_ROW_CODES: tuple[str, ...] = ("201", "202", "203", "204")
TOTAL_VALUE_ADDED_ROW = "209"
TOTAL_INTERMEDIATE_CODE = "190"
TOTAL_INPUT_ROW = "210"
TOTAL_ROW_CODES: tuple[str, ...] = (TOTAL_INTERMEDIATE_CODE, TOTAL_VALUE_ADDED_ROW, TOTAL_INPUT_ROW)

FINAL_DEMAND_COLUMN_CODES: tuple[str, ...] = ("301", "302", "303", "304", "305", "306")
IMPORT_COLUMN_CODES: tuple[str, ...] = ("401", "402", "403", "404")
STOCK_COLUMN_CODES: tuple[str, ...] = ("501", "502", "503")
# Subtotals must never be summed alongside their components.
SUBTOTAL_COLUMN_CODES: tuple[str, ...] = ("190", "309", "409", "509", "310", "700")
GROSS_OUTPUT_COLUMN = "600"

VALUATION_BASIS = {
    "price_basis": "purchasers_prices",
    "import_treatment": "import_inclusive",
    "import_specific_coefficients_available": True,
    "import_column_sign": "negative_in_demand_columns_positive_in_cell_import_field",
    "margins_itemised": ["wholesale", "retail", "transport"],
    "producers_price_derivable": True,
    "producers_price_formula": "PURCHASER - WHOLESALE - RETAIL - TRANSPORT",
    "domestic_only_derivable": True,
    "domestic_only_formula": "PURCHASER - IMPORT",
    "unit": "THB_million_or_thousand_as_published",
    "unit_note": (
        "The published file states no unit in the data sheet. Technical "
        "coefficients are RATIOS and are therefore unit-free, so C3 does not "
        "depend on resolving the level unit; absolute levels are not reported "
        "as monetary amounts anywhere in C3."
    ),
}

NESDC_ROOT = "https://www.nesdc.go.th"
NESDC_SITEMAP_INDEX = f"{NESDC_ROOT}/sitemap.xml"
NESDC_IO_LEGACY_PAGE = f"{NESDC_ROOT}/main.php?filename=io_page"
# C3-R1: the WordPress MEDIA LIBRARY. The aggregated workbooks (x58/x26/x16)
# live here and are NOT linked from any /download/ page, which is exactly why
# C3's download-page-only discovery missed them.
NESDC_MEDIA_API = f"{NESDC_ROOT}/wp-json/wp/v2/media"

# Published aggregation levels of the same 2015 table.
AGGREGATION_LEVELS = {
    "x58": {"sector_count": 58, "search": "DataIO2015x58"},
    "x26": {"sector_count": 26, "search": "DataIO2015x26"},
    "x16": {"sector_count": 16, "search": "DataIO2015x16"},
}

_DOWNLOAD_SLUG = re.compile(
    r"<loc>(https://www\.nesdc\.go\.th/(?:en/)?download/[^<]*"
    r"(?:i-o-table|input-output)[^<]*)</loc>",
    re.IGNORECASE,
)
_DDL_LINK = re.compile(
    r'href="(https://www\.nesdc\.go\.th/\?p=\d+&(?:amp;)?ddl=\d+)"[^>]*class="([^"]*)"',
    re.IGNORECASE,
)
_USER_AGENT = "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 Chrome/120.0"


class IOSourceError(RuntimeError):
    """An official I/O source could not be discovered or retrieved."""


class IOParseError(ValueError):
    """The workbook does not match the structure C3 requires."""


class WorkbookIdentityError(IOParseError):
    """A downloaded file is not the workbook it claims to be.

    Catches the two failure modes that look like success: an HTML error page
    saved under an .xlsx name, and a real workbook whose internal title belongs
    to a different table or aggregation level.
    """


class WafChallengeError(IOSourceError):
    """The WAF challenge did not clear within the allowed retries."""


# The WAF answers a first request with a Set-Cookie and a 302 back to the SAME
# URL. urllib's redirect handler treats that as an infinite loop, and a handler
# that rebuilds the request as `unverifiable` gets its cookies withheld by the
# cookie policy. Retrying explicitly is simpler and provably correct: the jar is
# populated by the first attempt, so the second request carries the cookie.
WAF_MAX_ATTEMPTS = 3


def fetch(
    url: str,
    timeout: int = 180,
    opener=None,
    attempts: int = WAF_MAX_ATTEMPTS,
    accept_statuses: tuple[int, ...] = (200,),
):
    """Retrieve a URL, returning payload and HTTP metadata.

    Bounded retry on the WAF's self-redirect. The bound matters: an unbounded
    retry would spin forever against a genuine redirect loop.

    ``accept_statuses`` exists because NESDC's page router answers its own
    content pages with HTTP 404 while serving a full, correct body. That is
    accepted ONLY where the caller opts in — a 404 on a FILE download stays
    fatal, so a missing workbook can never be mistaken for a served one.
    """
    opener = opener or urllib.request
    last_error: Exception | None = None
    for attempt in range(1, attempts + 1):
        request = urllib.request.Request(url, headers={"User-Agent": _USER_AGENT})
        try:
            with opener.open(request, timeout=timeout) as response:
                payload = response.read()
                return payload, {
                    "requested_url": url,
                    "final_url": response.geturl(),
                    "http_status": response.status,
                    "http_content_type": response.headers.get("Content-Type"),
                    "http_last_modified": response.headers.get("Last-Modified"),
                    "content_length": len(payload),
                    "waf_attempts": attempt,
                }
        except urllib.error.HTTPError as exc:
            if exc.code in accept_statuses and exc.code != 302:
                payload = exc.read()
                return payload, {
                    "requested_url": url,
                    "final_url": exc.url,
                    "http_status": exc.code,
                    "http_content_type": exc.headers.get("Content-Type"),
                    "http_last_modified": exc.headers.get("Last-Modified"),
                    "content_length": len(payload),
                    "waf_attempts": attempt,
                    "status_accepted_by_caller": True,
                }
            # 302 back to the same URL is the challenge; anything else is real.
            if exc.code != 302:
                raise
            last_error = exc
    raise WafChallengeError(
        f"{url}: the WAF still returned 302 after {attempts} attempts. "
        f"Last error: {last_error}"
    )


def discover_io_downloads(sitemap_xml: str) -> list[str]:
    """Return official I/O download-page URLs found in NESDC's own sitemap.

    Discovery starts from the sitemap rather than a remembered file URL because
    NESDC reorganised their site: the historical I/O page URL now serves an
    unrelated WordPress page with no I/O links on it at all.
    """
    return sorted(set(_DOWNLOAD_SLUG.findall(sitemap_xml)))


def resolve_download_url(page_html: str) -> list[dict]:
    """Extract the file-download links from an official download page.

    The pages expose files through a ``?p=<post>&ddl=<file>`` redirect rather
    than a direct href, with the file type carried in the anchor's CSS class.
    """
    found = []
    for href, css in _DDL_LINK.findall(page_html):
        kind = "unknown"
        for candidate in ("excel", "pdf", "word", "zip"):
            if candidate in css.lower():
                kind = candidate
                break
        found.append({"url": href.replace("&amp;", "&"), "file_kind": kind})
    return found


# --- workbook parsing --------------------------------------------------------


@dataclass
class IOTable:
    """One parsed official I/O table."""

    table_year: str
    source_url: str
    file_path: str
    sha256: str
    sheet_name: str
    sector_codes: list[str]
    # cells[(row_code, col_code)] -> {purchaser, wholesale, retail, transport, import}
    cells: dict[tuple[str, str], dict[str, float]]
    gross_output: dict[str, float]
    value_added: dict[str, float]
    diagnostics: dict = field(default_factory=dict)

    def z(self, row_code: str, col_code: str, measure: str = "purchaser") -> float:
        """One intermediate cell. Absent cells are structural zeros, not gaps."""
        cell = self.cells.get((row_code, col_code))
        return cell[measure] if cell else 0.0


def parse_io_workbook(
    path: Path, table_year: str, source_url: str, sha256: str,
    sector_count: int = SECTOR_CODE_COUNT,
) -> IOTable:
    """Parse the published long-format workbook into typed blocks.

    Labels and totals are validated; worksheet coordinates are never assumed.
    ``sector_count`` selects the aggregation level — NESDC publishes the same
    2015 table at 180, 58, 26 and 16 sectors with an identical layout.
    """
    sector_codes = tuple(f"{i:03d}" for i in range(1, sector_count + 1))
    import openpyxl

    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        if len(workbook.worksheets) != 1:
            raise IOParseError(
                f"Expected a single data sheet; found {len(workbook.worksheets)}: "
                f"{[w.title for w in workbook.worksheets]}"
            )
        sheet = workbook.worksheets[0]
        rows = list(sheet.iter_rows(values_only=True))
    finally:
        workbook.close()

    # Locate the header by LABEL, not by row number.
    header_index = None
    for index, row in enumerate(rows[:20]):
        values = [str(v).strip().upper() if v is not None else "" for v in row]
        if values[:2] == ["ROW", "COLUMN"] and "PURCHASER" in values:
            header_index = index
            break
    if header_index is None:
        raise IOParseError(
            "No 'ROW | COLUMN | PURCHASER ...' header row found; the workbook "
            "layout is not the one C3 was validated against."
        )
    header = [str(v).strip().upper() if v is not None else "" for v in rows[header_index]]
    expected = ["ROW", "COLUMN", "PURCHASER", "WHOLESALE", "RETAIL", "TRANSPORT", "IMPORT"]
    if header[: len(expected)] != expected:
        raise IOParseError(f"Unexpected header {header!r}; expected {expected!r}.")

    measures = ("purchaser", "wholesale", "retail", "transport", "import")
    cells: dict[tuple[str, str], dict[str, float]] = {}
    duplicates: list[tuple[str, str]] = []
    for row in rows[header_index + 1:]:
        if row[0] is None or row[1] is None:
            continue
        key = (str(row[0]).strip(), str(row[1]).strip())
        if key in cells:
            duplicates.append(key)
            continue
        cells[key] = {
            name: float(row[2 + offset]) if row[2 + offset] is not None else 0.0
            for offset, name in enumerate(measures)
        }
    if duplicates:
        raise IOParseError(
            f"{len(duplicates)} duplicate (ROW, COLUMN) keys, e.g. {duplicates[:3]}. "
            "A duplicated cell would silently double-count an input."
        )

    codes_seen = {key[0] for key in cells} | {key[1] for key in cells}
    unknown = sorted(
        code for code in codes_seen
        if code not in sector_codes
        and code not in VALUE_ADDED_ROW_CODES + TOTAL_ROW_CODES
        + FINAL_DEMAND_COLUMN_CODES + IMPORT_COLUMN_CODES + STOCK_COLUMN_CODES
        + SUBTOTAL_COLUMN_CODES + (GROSS_OUTPUT_COLUMN,)
    )
    if unknown:
        raise IOParseError(
            f"Codes present in the workbook but not in C3's documented code "
            f"universe: {unknown}. Refusing to guess whether they are sectors, "
            "totals, or demand columns."
        )

    gross_output = {
        code: (cells.get((code, GROSS_OUTPUT_COLUMN)) or {}).get("purchaser", 0.0)
        for code in sector_codes
    }
    value_added = {
        code: (cells.get((TOTAL_VALUE_ADDED_ROW, code)) or {}).get("purchaser", 0.0)
        for code in sector_codes
    }

    # --- accounting validation, reported not assumed -------------------------
    total_input_mismatches = [
        code for code in sector_codes
        if abs(
            gross_output[code]
            - (cells.get((TOTAL_INPUT_ROW, code)) or {}).get("purchaser", 0.0)
        ) > max(1.0, 1e-9 * abs(gross_output[code]))
    ]
    column_residuals = {}
    for code in sector_codes:
        intermediate = sum(
            (cells.get((row_code, code)) or {}).get("purchaser", 0.0)
            for row_code in sector_codes
        )
        total = intermediate + value_added[code]
        column_residuals[code] = abs(total - gross_output[code]) / max(
            abs(gross_output[code]), 1.0
        )

    zero_output = sorted(code for code in sector_codes if gross_output[code] <= 0)
    absent = sorted(
        code for code in sector_codes
        if code not in {k[0] for k in cells} and code not in {k[1] for k in cells}
    )

    diagnostics = {
        "sheet_name": sheet.title,
        "header_row_index": header_index + 1,
        "header": header,
        "data_cells": len(cells),
        "sector_codes_expected": sector_count,
        "sector_codes_present_as_row": len({k[0] for k in cells} & set(sector_codes)),
        "sector_codes_present_as_column": len({k[1] for k in cells} & set(sector_codes)),
        "sector_codes_absent_entirely": absent,
        "duplicate_cell_keys": len(duplicates),
        "gross_output_total": sum(gross_output.values()),
        "value_added_total": sum(value_added.values()),
        "intermediate_total": sum(
            (cells.get((r, c)) or {}).get("purchaser", 0.0)
            for r in sector_codes for c in sector_codes
        ),
        "import_content_total": sum(
            (cells.get((r, c)) or {}).get("import", 0.0)
            for r in sector_codes for c in sector_codes
        ),
        "gross_output_equals_total_input_row": not total_input_mismatches,
        "gross_output_mismatch_sectors": total_input_mismatches,
        "max_column_identity_residual": max(column_residuals.values()),
        "column_identity_worst_sector": max(column_residuals, key=column_residuals.get),
        "zero_or_negative_gross_output_sectors": zero_output,
        "valuation": dict(VALUATION_BASIS),
    }

    return IOTable(
        table_year=table_year,
        source_url=source_url,
        file_path=str(path).replace("\\", "/"),
        sha256=sha256,
        sheet_name=sheet.title,
        sector_codes=list(sector_codes),
        cells=cells,
        gross_output=gross_output,
        value_added=value_added,
        diagnostics=diagnostics,
    )


_DEFINITION = re.compile(r"^\s*(\d{3})\s+([A-Z][^\n]{2,90})$", re.M)


def parse_sector_definitions(pdf_path: Path) -> dict[str, dict[str, str]]:
    """Extract sector code -> {label, definition} from the official description PDF.

    The workbook carries codes only, so every human-readable label and every
    crosswalk justification in C3 comes from this document.
    """
    from pypdf import PdfReader

    reader = PdfReader(pdf_path)
    text = "\n".join(page.extract_text() for page in reader.pages)
    text = re.sub(r"\(\d+\)\s*\n", "", text)

    definitions: dict[str, dict[str, str]] = {}
    matches = list(_DEFINITION.finditer(text))
    for index, match in enumerate(matches):
        code = match.group(1)
        if code in definitions or code not in SECTOR_CODES:
            continue
        end = matches[index + 1].start() if index + 1 < len(matches) else len(text)
        definitions[code] = {
            "label": " ".join(match.group(2).split()),
            "definition": " ".join(text[match.end():end].split()),
        }
    missing = [code for code in SECTOR_CODES if code not in definitions]
    if missing:
        raise IOParseError(
            f"{len(missing)} sector codes have no label in the classification "
            f"document, e.g. {missing[:5]}. C3 will not invent labels."
        )
    return definitions


def discover_media_files(payload: bytes) -> list[dict]:
    """Parse NESDC's WordPress media REST response into file records.

    This is official first-party discovery: the publisher's own API listing its
    own uploads. C3 never queried it, which is why the 58-sector workbooks were
    invisible to it.
    """
    import json as _json

    try:
        items = _json.loads(payload.decode("utf-8", "replace"))
    except ValueError as exc:
        raise IOSourceError(f"Media API did not return JSON: {exc}") from exc
    if not isinstance(items, list):
        return []
    found = []
    for item in items:
        url = item.get("source_url")
        if not url:
            continue
        title = item.get("title")
        if isinstance(title, dict):
            title = title.get("rendered", "")
        found.append(
            {
                "url": url,
                "title": str(title or ""),
                "filename": url.rsplit("/", 1)[-1],
                "media_id": item.get("id"),
                "date": item.get("date"),
            }
        )
    return found


def validate_workbook_identity(
    path: Path, expected_title_fragment: str, expected_sheet: str | None = None
) -> dict:
    """Confirm a downloaded workbook really is the table it claims to be.

    Checks the ZIP magic (an HTML error page saved as .xlsx fails here), the
    sheet name, and the workbook's own internal title cell.
    """
    import openpyxl

    head = path.read_bytes()[:4]
    if head[:2] != b"PK":
        raise WorkbookIdentityError(
            f"{path.name} is not a ZIP-based workbook (leading bytes {head!r}); "
            "an HTML error page saved under an .xlsx name would look like this."
        )

    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        sheets = [sheet.title for sheet in workbook.worksheets]
        first = workbook.worksheets[0]
        rows = []
        for index, row in enumerate(first.iter_rows(values_only=True)):
            rows.append(row)
            if index >= 2:
                break
    finally:
        workbook.close()

    internal_title = str(rows[0][0]).strip() if rows and rows[0] and rows[0][0] else ""
    if expected_title_fragment.lower() not in internal_title.lower():
        raise WorkbookIdentityError(
            f"{path.name} internal title {internal_title!r} does not contain "
            f"{expected_title_fragment!r}. Refusing to treat it as that table."
        )
    if expected_sheet is not None and expected_sheet not in sheets:
        raise WorkbookIdentityError(
            f"{path.name} has sheets {sheets}, expected {expected_sheet!r}."
        )
    return {
        "sheets": sheets,
        "internal_title": internal_title,
        "header": list(rows[1]) if len(rows) > 1 else [],
    }
