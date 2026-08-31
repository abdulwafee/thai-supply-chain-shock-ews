"""Guarded MPI workbook reader and revision taxonomy (Task D2).

The guard
---------
Task D2 audits *when* OIE published each reference month. It must not become a
back door for reading the months that the locked evaluation still depends on.
So this module separates two things that are easy to conflate:

* **Structural metadata** -- which month columns exist, what they are labelled,
  which carry a marker glyph, the file checksum, the HTTP headers. None of this
  is an index value. It stays readable for every month, which is precisely what
  a publication-timing audit needs.
* **Numerical index values** -- readable only at or before
  ``numerical_audit_cutoff_month``. Beyond it :class:`GuardedMonthlyReader`
  raises :class:`NumericalCutoffError` instead of returning a number.

Downloading a current official file for provenance is not permission to parse
every numerical month inside it. The guard is therefore enforced at the point of
extraction, not by asking callers to remember a rule.

The revision taxonomy
---------------------
Two editions of the same reference month can differ for very different reasons,
and collapsing them into one "revision" bucket destroys the distinction that
matters. :func:`classify_revision` keeps them apart, and refuses to call a
difference an ordinary revision when the two editions are not comparable in the
first place (different base year, different classification scheme).
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path

import yaml

__all__ = [
    "REVISION_CLASSES",
    "THAI_MONTH_ABBREVIATIONS",
    "GuardedMonthlyReader",
    "MonthColumn",
    "NumericalCutoffError",
    "classify_revision",
    "default_timing_config_path",
    "load_timing_config",
    "parse_thai_month_header",
]

# Abbreviated Thai month names as they appear in the OIE workbook header row.
THAI_MONTH_ABBREVIATIONS = {
    "ม.ค.": 1, "ก.พ.": 2, "มี.ค.": 3, "เม.ย.": 4, "พ.ค.": 5, "มิ.ย.": 6,
    "ก.ค.": 7, "ส.ค.": 8, "ก.ย.": 9, "ต.ค.": 10, "พ.ย.": 11, "ธ.ค.": 12,
}

REVISION_CLASSES = (
    "rounding_only",
    "ordinary_revision",
    "preliminary_to_later",
    "annual_revision",
    "base_year_rebasing",
    "classification_change",
    "industry_definition_change",
    "not_comparable",
)

_BUDDHIST_ERA_OFFSET = 543
_MONTHLY_SHEET = "รายเดือน"
_DIVISION_PATTERN = re.compile(r"\s*(TSIC|ISIC)\s*:\s*(\d{2})(?:\s|$)")
_BE_YEAR_PATTERN = re.compile(r"^25\d\d$")


class NumericalCutoffError(RuntimeError):
    """Raised when numerical extraction is attempted beyond the audit cutoff."""


def default_timing_config_path() -> Path:
    return Path(__file__).resolve().parents[3] / "configs" / "oie_mpi_timing.yaml"


def load_timing_config(path=None) -> dict:
    target = Path(path) if path is not None else default_timing_config_path()
    with open(target, encoding="utf-8") as handle:
        return yaml.safe_load(handle)


def parse_thai_month_header(year_cell, month_cell, carried_year):
    """Resolve one workbook header column to ``(reference_month, marker, year)``.

    The workbook writes the Buddhist-Era year once per 12-column block and
    leaves the rest blank, so the year is carried forward. A month label may
    carry a trailing marker glyph (observed: ``มิ.ย.*``); the marker is returned
    separately rather than silently stripped, because whether a marker means
    "preliminary" is a question about OIE's documentation, not about parsing.
    """
    year = carried_year
    if year_cell is not None and _BE_YEAR_PATTERN.match(str(year_cell).strip()):
        year = int(str(year_cell).strip()) - _BUDDHIST_ERA_OFFSET

    raw = "" if month_cell is None else str(month_cell).strip()
    if not raw:
        return None, None, year
    bare = raw.rstrip("*").strip()
    month = THAI_MONTH_ABBREVIATIONS.get(bare)
    if month is None or year is None:
        return None, None, year
    marker = raw[len(bare):].strip() or None
    return f"{year:04d}-{month:02d}", marker, year


@dataclass(frozen=True)
class MonthColumn:
    """One month column's *structural* description. Carries no index value."""

    column_index: int
    reference_month: str
    marker: str = None

    @property
    def has_marker(self) -> bool:
        return self.marker is not None


@dataclass
class GuardedMonthlyReader:
    """Reads an OIE monthly workbook under a hard numerical cutoff.

    ``structure()`` describes every month column, including months beyond the
    cutoff. ``division_series()`` returns values only at or before the cutoff.
    ``require_readable()`` converts a rule into an exception.
    """

    path: Path
    cutoff_month: str
    sheet_name: str = _MONTHLY_SHEET
    _grid: list = field(default_factory=list, repr=False)
    _columns: list = field(default_factory=list, repr=False)
    _loaded: bool = field(default=False, repr=False)

    @classmethod
    def from_config(cls, path, config=None, **kwargs):
        cfg = config if config is not None else load_timing_config()
        return cls(path=Path(path), cutoff_month=cfg["numerical_audit_cutoff_month"], **kwargs)

    # -- guard ------------------------------------------------------------
    def is_readable(self, reference_month: str) -> bool:
        return reference_month <= self.cutoff_month

    def require_readable(self, reference_month: str) -> None:
        """Raise unless numerical extraction is permitted for this month."""
        if not self.is_readable(reference_month):
            raise NumericalCutoffError(
                f"numerical extraction refused for reference_month={reference_month}: beyond "
                f"numerical_audit_cutoff_month={self.cutoff_month}. Structural "
                "metadata is available "
                "via structure(); index values are not."
            )

    # -- loading ----------------------------------------------------------
    def _load(self) -> None:
        if self._loaded:
            return
        from openpyxl import load_workbook

        workbook = load_workbook(self.path, read_only=True, data_only=True)
        try:
            sheet = workbook[self.sheet_name]
            self._grid = [tuple(row) for row in sheet.iter_rows(values_only=True)]
        finally:
            workbook.close()

        year_row = self._grid[3] if len(self._grid) > 3 else ()
        month_row = self._grid[4] if len(self._grid) > 4 else ()
        columns, carried = [], None
        for index in range(max(len(year_row), len(month_row))):
            year_cell = year_row[index] if index < len(year_row) else None
            month_cell = month_row[index] if index < len(month_row) else None
            reference_month, marker, carried = parse_thai_month_header(
                year_cell, month_cell, carried
            )
            if reference_month is not None:
                columns.append(MonthColumn(index, reference_month, marker))
        self._columns = columns
        self._loaded = True

    # -- structural (always permitted) ------------------------------------
    def structure(self) -> list:
        self._load()
        return list(self._columns)

    def coverage(self):
        months = sorted(column.reference_month for column in self.structure())
        return (months[0], months[-1]) if months else (None, None)

    def marked_months(self) -> list:
        return [c.reference_month for c in self.structure() if c.has_marker]

    def base_year(self):
        """Buddhist-Era base year declared in the sheet's own header note."""
        self._load()
        for row in self._grid[:6]:
            text = str(row[0] or "")
            if "ฐานเฉลี่ยรายเดือน" in text:
                found = re.search(r"(25\d\d)", text)
                if found:
                    return int(found.group(1)) - _BUDDHIST_ERA_OFFSET
        return None

    def classification_scheme(self):
        """``TSIC`` or ``ISIC`` -- whichever the division rows actually use."""
        self._load()
        schemes = set()
        for row in self._grid[5:]:
            match = _DIVISION_PATTERN.match(str(row[2] or "")) if len(row) > 2 else None
            if match:
                schemes.add(match.group(1))
        if len(schemes) == 1:
            return schemes.pop()
        return sorted(schemes) or None

    def divisions(self) -> list:
        self._load()
        seen = []
        for row in self._grid[5:]:
            match = _DIVISION_PATTERN.match(str(row[2] or "")) if len(row) > 2 else None
            if match:
                code = int(match.group(2))
                if code not in seen:
                    seen.append(code)
        return sorted(seen)

    # -- numerical (guarded) ----------------------------------------------
    def division_series(self, division: int, allow_beyond_cutoff: bool = False) -> dict:
        """Monthly values for one division, truncated at the cutoff.

        Months beyond the cutoff are omitted, not returned as null -- omission
        records "not read", whereas null would assert "read and empty".
        ``allow_beyond_cutoff`` exists only so tests can prove the guard fires;
        it raises rather than widening the window.
        """
        self._load()
        row = self._division_row(division)
        if row is None:
            return {}
        series = {}
        for column in self._columns:
            if not self.is_readable(column.reference_month):
                if allow_beyond_cutoff:
                    self.require_readable(column.reference_month)
                continue
            value = row[column.column_index] if column.column_index < len(row) else None
            if isinstance(value, (int, float)) and not isinstance(value, bool):
                series[column.reference_month] = float(value)
        return series

    def value_at(self, division: int, reference_month: str):
        """Single guarded lookup. Raises beyond the cutoff."""
        self.require_readable(reference_month)
        return self.division_series(division).get(reference_month)

    def _division_row(self, division: int):
        for row in self._grid[5:]:
            match = _DIVISION_PATTERN.match(str(row[2] or "")) if len(row) > 2 else None
            if match and int(match.group(2)) == division:
                return row
        return None


def classify_revision(
    earlier,
    later,
    earlier_base_year=None,
    later_base_year=None,
    earlier_scheme=None,
    later_scheme=None,
    earlier_divisions=None,
    later_divisions=None,
    displayed_decimals=None,
    earlier_was_marked=False,
    crosses_annual_revision=False,
):
    """Classify why two editions of one reference month differ.

    Comparability is decided BEFORE magnitude. A base-year change or a
    classification change makes the two numbers different quantities, so the
    size of their difference is not evidence of a revision and must not be
    reported as one.
    """
    if earlier_scheme and later_scheme and earlier_scheme != later_scheme:
        return "classification_change"
    if (
        earlier_base_year is not None
        and later_base_year is not None
        and earlier_base_year != later_base_year
    ):
        return "base_year_rebasing"
    if earlier_divisions is not None and later_divisions is not None:
        if set(earlier_divisions) != set(later_divisions):
            return "industry_definition_change"
    if earlier is None or later is None:
        return "not_comparable"

    if displayed_decimals is not None:
        if round(earlier, displayed_decimals) == round(later, displayed_decimals):
            return "rounding_only"
    if earlier == later:
        return "rounding_only"
    if earlier_was_marked:
        return "preliminary_to_later"
    if crosses_annual_revision:
        return "annual_revision"
    return "ordinary_revision"
